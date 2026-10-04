"""Genere les slides de soutenance pour EduPaie."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


# Couleurs de la charte
BLEU = RGBColor(0x1E, 0x3A, 0x8A)
BLEU_CLAIR = RGBColor(0x25, 0x63, 0xEB)
VERT = RGBColor(0x10, 0xB9, 0x81)
ORANGE = RGBColor(0xF5, 0x9E, 0x0B)
ROUGE = RGBColor(0xEF, 0x44, 0x44)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
GRIS = RGBColor(0x64, 0x74, 0x8B)


def ajouter_slide(prs, titre, sous_titre=None, fond=BLANC, couleur_titre=BLEU):
    """Ajoute une slide avec titre."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Fond
    fond_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    fond_shape.fill.solid()
    fond_shape.fill.fore_color.rgb = fond
    fond_shape.line.fill.background()

    # Titre
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(1))
    tf = tb.text_frame
    tf.text = titre
    p = tf.paragraphs[0]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = couleur_titre
    p.font.name = "Calibri"

    # Sous-titre
    if sous_titre:
        tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11), Inches(0.6))
        tf2 = tb2.text_frame
        tf2.text = sous_titre
        p2 = tf2.paragraphs[0]
        p2.font.size = Pt(16)
        p2.font.color.rgb = GRIS
        p2.font.name = "Calibri"

    return slide


def ajouter_texte(slide, gauche, haut, largeur, hauteur, lignes, taille=18,
                  couleur=RGBColor(0x1E, 0x29, 0x3B), gras=False):
    """Ajoute un bloc de texte."""
    tb = slide.shapes.add_textbox(Inches(gauche), Inches(haut), Inches(largeur), Inches(hauteur))
    tf = tb.text_frame
    tf.word_wrap = True

    for i, ligne in enumerate(lignes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ligne
        p.font.size = Pt(taille)
        p.font.color.rgb = couleur
        p.font.bold = gras
        p.font.name = "Calibri"
        p.space_after = Pt(8)

    return tb


def ajouter_carte(slide, x, y, largeur, hauteur, titre, valeur, couleur):
    """Ajoute une carte coloree (comme dans l'app)."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(largeur), Inches(hauteur)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLANC
    shape.line.color.rgb = couleur
    shape.line.width = Pt(2)

    tf = shape.text_frame
    tf.text = titre + "\n" + valeur
    p1 = tf.paragraphs[0]
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = GRIS
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = valeur
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = couleur
    p2.alignment = PP_ALIGN.CENTER


# ============================================================
# GENERATION
# ============================================================

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)


# ---------- SLIDE 1 : TITRE ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = BLEU
bg.line.fill.background()

tb = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(2))
tf = tb.text_frame
tf.text = "EduPaie"
p = tf.paragraphs[0]
p.font.size = Pt(72)
p.font.bold = True
p.font.color.rgb = BLANC
p.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = "Gestion des paiements scolaires"
p2.font.size = Pt(28)
p2.font.color.rgb = RGBColor(0xDB, 0xEA, 0xFE)
p2.alignment = PP_ALIGN.CENTER

p3 = tf.add_paragraph()
p3.text = ""
p3.font.size = Pt(14)

p4 = tf.add_paragraph()
p4.text = "Francoise97  |  04 octobre 2026"
p4.font.size = Pt(18)
p4.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
p4.alignment = PP_ALIGN.CENTER


# ---------- SLIDE 2 : CONTEXTE ----------
slide = ajouter_slide(prs, "Contexte", "Un probleme reel dans les etablissements scolaires")

ajouter_texte(slide, 0.8, 2.2, 11, 4, [
    "❌  Suivi manuel des paiements (cahier, Excel)",
    "",
    "❌  Aucune vue fiable sur qui a paye quoi",
    "",
    "❌  Reçus non numerotes, impossibles a retrouver en cas de litige",
    "",
    "❌  Aucun calcul automatique du solde restant du",
], taille=18)


# ---------- SLIDE 3 : MISSION ----------
slide = ajouter_slide(prs, "Mission", "Automatiser la gestion des paiements scolaires")

ajouter_texte(slide, 0.8, 2.2, 11, 4, [
    "✅  Enregistrer chaque eleve et ses frais",
    "",
    "✅  Enregistrer chaque paiement avec validation",
    "",
    "✅  Calculer automatiquement le solde restant du",
    "",
    "✅  Generer un recu numerote unique a chaque versement",
    "",
    "✅  Offrir une interface simple a utiliser (secretaire)",
], taille=18, couleur=RGBColor(0x1E, 0x29, 0x3B))


# ---------- SLIDE 4 : FONCTIONNALITES ----------
slide = ajouter_slide(prs, "Fonctionnalites", "6 ecrans principaux")

# Cartes 3x2
cartes = [
    (0.8, 2.2, "TABLEAU DE BORD", "Statistiques temps reel", BLEU),
    (4.9, 2.2, "ELEVES", "CRUD + recherche + filtres", VERT),
    (9.0, 2.2, "FICHE ELEVE", "Historique + situation", ORANGE),
    (0.8, 4.2, "PAIEMENTS", "Enregistrement + validation", VERT),
    (4.9, 4.2, "RECUS", "Liste + generation PDF", BLEU_CLAIR),
    (9.0, 4.2, "PARAMETRES", "Config centralisee", ROUGE),
]

for x, y, titre, desc, couleur in cartes:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(3.5), Inches(1.6)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLANC
    shape.line.color.rgb = couleur
    shape.line.width = Pt(2)

    tf = shape.text_frame
    tf.text = titre
    p = tf.paragraphs[0]
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = couleur
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = GRIS
    p2.alignment = PP_ALIGN.CENTER


# ---------- SLIDE 5 : ARCHITECTURE ----------
slide = ajouter_slide(prs, "Architecture", "3 couches separees")

# Couche UI
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.4))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0xDB, 0xEA, 0xFE)
shape.line.color.rgb = BLEU
tf = shape.text_frame
tf.text = "INTERFACE (PySide6)"
p = tf.paragraphs[0]
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = BLEU
p.alignment = PP_ALIGN.CENTER
p2 = tf.add_paragraph()
p2.text = "widgets/ · dialogs/ · main_window.py"
p2.font.size = Pt(12)
p2.font.color.rgb = GRIS
p2.alignment = PP_ALIGN.CENTER

# Couche Services
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(0.8), Inches(3.6), Inches(11.5), Inches(1.4))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0xD1, 0xFA, 0xE5)
shape.line.color.rgb = VERT
tf = shape.text_frame
tf.text = "LOGIQUE METIER (Services)"
p = tf.paragraphs[0]
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = VERT
p.alignment = PP_ALIGN.CENTER
p2 = tf.add_paragraph()
p2.text = "calcul du solde · validation · numerotation recus · generation PDF"
p2.font.size = Pt(12)
p2.font.color.rgb = GRIS
p2.alignment = PP_ALIGN.CENTER

# Couche Repositories
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(0.8), Inches(5.2), Inches(11.5), Inches(1.4))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0xFE, 0xE2, 0xE2)
shape.line.color.rgb = ROUGE
tf = shape.text_frame
tf.text = "ACCES DONNEES (Repositories)"
p = tf.paragraphs[0]
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = ROUGE
p.alignment = PP_ALIGN.CENTER
p2 = tf.add_paragraph()
p2.text = "SQL encapsule · aucune requete dans l'UI"
p2.font.size = Pt(12)
p2.font.color.rgb = GRIS
p2.alignment = PP_ALIGN.CENTER


# ---------- SLIDE 6 : BASE DE DONNEES ----------
slide = ajouter_slide(prs, "Base de donnees", "SQLite - 3 tables + 1 vue")

ajouter_texte(slide, 0.8, 2.0, 11, 5, [
    "📋  Table eleve : id, nom, prenom, classe, annee_scolaire, frais_totaux, telephone, nom_parent",
    "",
    "📋  Table recu : id, numero_unique (REC-AAAA-NNN), annee, sequence, montant_paye, solde_apres",
    "",
    "📋  Table paiement : id, eleve_id, recu_id, montant, date_paiement, mode_paiement, reference, observation",
    "",
    "🔍  Vue vue_solde_eleve : calcule automatiquement total_paye, solde_restant et statut",
    "",
    "🔒  Contraintes : CHECK (montant > 0), UNIQUE(annee, sequence), FK avec ON DELETE CASCADE",
], taille=14)


# ---------- SLIDE 7 : REGLES METIER ----------
slide = ajouter_slide(prs, "Regles metier", "Garanties et validations")

ajouter_texte(slide, 0.8, 2.0, 11, 5, [
    "💰  Calcul du solde  :  solde = frais_totaux - somme(paiements)",
    "",
    "🚫  Le solde ne peut JAMAIS etre negatif (validation avant enregistrement)",
    "",
    "📄  Numerotation unique des recus : REC-{ANNEE}-{NUMERO:03d}",
    "",
    "   → La sequence repart a 001 chaque annee",
    "",
    "✅  Transaction atomique : recu + paiement crees ensemble (rollback si erreur)",
    "",
    "🎯  Statuts : Non paye (0) / Partiellement paye / Solde",
], taille=15)


# ---------- SLIDE 8 : CHOIX TECHNIQUES ----------
slide = ajouter_slide(prs, "Choix techniques", "Pourquoi ces technologies ?")

# 3 colonnes
cols = [
    (0.8, "SQLite", [
        "Aucun serveur a installer",
        "Fichier unique portable",
        "Module sqlite3 standard",
        "Parfait pour du mono-poste",
    ], BLEU),
    (4.9, "PySide6", [
        "Framework Qt moderne",
        "Signaux / slots clairs",
        "Look natif Windows",
        "Licence LGPL",
    ], VERT),
    (9.0, "reportlab", [
        "PDF de qualite pro",
        "Controle total mise en page",
        "Licence BSD",
        "Reference du marche",
    ], ORANGE),
]

for x, titre, lignes, couleur in cols:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(2.0), Inches(3.5), Inches(4.3)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLANC
    shape.line.color.rgb = couleur
    shape.line.width = Pt(2)

    tf = shape.text_frame
    tf.text = titre
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = couleur
    p.alignment = PP_ALIGN.CENTER

    for ligne in lignes:
        p2 = tf.add_paragraph()
        p2.text = "• " + ligne
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)


# ---------- SLIDE 9 : DEMO ----------
slide = ajouter_slide(prs, "Demonstration live", "Passons a la pratique")

ajouter_texte(slide, 0.8, 3.0, 11, 2, [
    "🎬  Nous allons maintenant voir l'application en fonctionnement",
    "",
    "   → Ajout d'un eleve",
    "   → Enregistrement d'un paiement",
    "   → Generation du recu PDF",
], taille=20, couleur=BLEU)


# ---------- SLIDE 10 : CHIFFRES CLES ----------
slide = ajouter_slide(prs, "Chiffres cles", "Bilan du projet")

cartes = [
    (0.8, 2.2, "COMMITS GIT", "25", BLEU),
    (4.9, 2.2, "BRANCHES", "10", VERT),
    (9.0, 2.2, "LIGNES DE CODE", "~2500", ORANGE),
    (0.8, 4.2, "ECRANS", "6", BLEU_CLAIR),
    (4.9, 4.2, "ELEVES EN BASE", "18", VERT),
    (9.0, 4.2, "TESTS", "3", ROUGE),
]

for x, y, titre, valeur, couleur in cartes:
    ajouter_carte(slide, x, y, 3.5, 1.6, titre, valeur, couleur)


# ---------- SLIDE 11 : LIMITES ----------
slide = ajouter_slide(prs, "Limites & evolutions", "Ce qu'on pourrait ameliorer")

ajouter_texte(slide, 0.8, 2.0, 5.5, 5, [
    "⚠️  Limites actuelles",
    "",
    "•  Application mono-poste (pas de reseau)",
    "•  Pas d'authentification utilisateurs",
    "•  Pas d'export Excel / CSV",
    "•  Pas de montant en lettres",
], taille=14)

ajouter_texte(slide, 6.8, 2.0, 5.5, 5, [
    "🚀  Evolutions possibles",
    "",
    "•  Multi-utilisateurs (caissier, directeur)",
    "•  Synchronisation cloud",
    "•  Envoi recus par email / SMS",
    "•  Graphiques statistiques",
], taille=14, couleur=VERT)


# ---------- SLIDE 12 : MERCI ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = BLEU
bg.line.fill.background()

tb = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(2.5))
tf = tb.text_frame
tf.text = "Merci de votre attention"
p = tf.paragraphs[0]
p.font.size = Pt(48)
p.font.bold = True
p.font.color.rgb = BLANC
p.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = ""
p2.font.size = Pt(14)

p3 = tf.add_paragraph()
p3.text = "Des questions ?"
p3.font.size = Pt(28)
p3.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
p3.alignment = PP_ALIGN.CENTER


# ============================================================
# SAUVEGARDE
# ============================================================
prs.save("docs/soutenance_EduPaie.pptx")
print("OK slides generees : docs/soutenance_EduPaie.pptx")
print("Total slides :", len(prs.slides))
